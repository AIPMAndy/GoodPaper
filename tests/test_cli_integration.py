"""Real command-line smoke tests using generated manuscripts and bundled demo templates."""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "template_packages" / "demo_template" / "demo_template.docx"


def run_cli(*args):
    return subprocess.run(
        [sys.executable, str(ROOT / "app.py"), *map(str, args)],
        cwd=ROOT, text=True, capture_output=True, timeout=35,
    )


def test_cli_check_and_format_keep_original_intact(tmp_path):
    manuscript = tmp_path / "paper.docx"
    output = tmp_path / "formatted.docx"
    doc = Document()
    doc.add_paragraph("Test paper heading")
    doc.add_paragraph("This is a first research paragraph.")
    doc.add_paragraph("References")
    doc.add_paragraph("[1] Example source.")
    doc.save(manuscript)
    before = manuscript.read_bytes()

    checked = run_cli("check", manuscript, "-t", TEMPLATE)
    assert checked.returncode == 0, checked.stderr
    report_path = tmp_path / "paper_report.json"
    assert report_path.exists()
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert isinstance(report.get("issues"), list)

    formatted = run_cli("format", manuscript, "-t", TEMPLATE, "-o", output)
    assert formatted.returncode == 0, formatted.stderr
    assert manuscript.read_bytes() == before
    assert output.exists()
    result = Document(output)
    assert "Test paper heading" in [p.text for p in result.paragraphs]


def test_cli_prevents_overwriting_manuscript(tmp_path):
    path = tmp_path / "paper.docx"
    d = Document()
    d.add_paragraph("Preserve this")
    d.save(path)
    before = path.read_bytes()
    result = run_cli("format", path, "-t", TEMPLATE, "-o", path)
    assert result.returncode != 0
    assert path.read_bytes() == before
